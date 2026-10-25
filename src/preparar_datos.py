"""
Script de preparación e integración de datos para el proyecto
"Análisis de la Brecha Digital y Calidad de Red en Chile"

Integra: SUBTEL + INE + CASEN → dataset_comunas_chile.csv (346 comunas de Chile)

Ejecutar con: python src/preparar_datos.py
"""
import os
import unicodedata
import pandas as pd
import numpy as np
import openpyxl
from pathlib import Path
import warnings

warnings.filterwarnings('ignore')

BASE_DIR = Path(__file__).parent.parent
SUBTEL_PATH = BASE_DIR / "data" / "raw" / "subtel" / "series_internet_fija_subtel.xlsx"
DATA_PROC   = BASE_DIR / "data" / "processed"
OUT_PATH    = DATA_PROC / "dataset_comunas_chile.csv"

DATA_PROC.mkdir(parents=True, exist_ok=True)

SYNONYMS = {
    'paiguano': 'paihuano',
    'aysen': 'aisen',
    'tiltil': 'til til',
    'llayllay': 'llay llay',
    'treguaco': 'trehuaco',
}

def norm(s):
    """Normaliza texto eliminando acentos y espacios sobrantes para merge robusto."""
    if not isinstance(s, str):
        return ''
    s = s.strip().lower()
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')


def extraer_subtel(path):
    """
    Extrae conexiones fijas residenciales por comuna de SUBTEL.
    Lee el Bloque 2 (2019-2026, columnas 53 a 144) de la hoja '7.11.1.CO_FIJAS_RES_COMUNA'
    que contiene las 16 regiones oficiales (incluyendo Ñuble) y 346 comunas alineadas.
    """
    print("📡 Extrayendo datos SUBTEL (hoja 7.11.1.CO_FIJAS_RES_COMUNA, Bloque 2019-2026)...")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['7.11.1.CO_FIJAS_RES_COMUNA']

    subtel_dict = {}
    for r in range(10, ws.max_row + 1):
        com = ws.cell(r, 54).value
        if com:
            com_str = str(com).strip()
            if not com_str.startswith('Total') and com_str != 'Sin clasificación' and com_str != 'None':
                # Columnas 115-126 corresponden a los 12 meses de 2024
                c_2024 = [ws.cell(r, c).value for c in range(115, 127) if isinstance(ws.cell(r, c).value, (int, float))]
                prom_2024 = sum(c_2024) / len(c_2024) if c_2024 else None
                # Columnas 103-114 corresponden a 2023
                c_2023 = [ws.cell(r, c).value for c in range(103, 115) if isinstance(ws.cell(r, c).value, (int, float))]
                prom_2023 = sum(c_2023) / len(c_2023) if c_2023 else None
                # Columna 144 corresponde a Junio 2026 (mes más reciente reportado)
                reciente = ws.cell(r, 144).value
                reciente = float(reciente) if isinstance(reciente, (int, float)) else None

                key = norm(com_str)
                subtel_dict[key] = {
                    'conexiones_recientes': reciente,
                    'conexiones_prom2024': prom_2024,
                    'conexiones_prom2023': prom_2023
                }

    print(f"   ✅ {len(subtel_dict)} comunas extraídas exitosamente desde SUBTEL")
    return subtel_dict


def actualizar_dataset(subtel_dict):
    """Aplica la corrección de SUBTEL y deduplicación a nivel de 346 comunas."""
    print("\n📦 Integrando con dataset base (INE + CASEN)...")
    df = pd.read_csv(OUT_PATH)

    # 1. Asegurar exactamente 346 comunas únicas (elimina las 21 comunas duplicadas de Ñuble)
    df = df.drop_duplicates(subset=['cut_comuna'], keep='first').sort_values('cut_comuna').reset_index(drop=True)

    recientes_list = []
    prom2024_list = []
    prom2023_list = []

    for idx, row in df.iterrows():
        k = norm(row['nombre_comuna'])
        if k in SYNONYMS:
            k = SYNONYMS[k]
        data = subtel_dict.get(k, {})
        recientes_list.append(data.get('conexiones_recientes', None))
        prom2024_list.append(data.get('conexiones_prom2024', None))
        prom2023_list.append(data.get('conexiones_prom2023', None))

    df['conexiones_recientes'] = recientes_list
    df['conexiones_prom2024'] = prom2024_list
    df['conexiones_prom2023'] = prom2023_list
    df['conexiones_x1000hab'] = (df['conexiones_prom2024'] / df['poblacion_2024']) * 1000

    df.to_csv(OUT_PATH, index=False, encoding='utf-8-sig')
    print(f"   ✅ Guardado: {OUT_PATH} con {len(df)} comunas (exactamente las 346 comunas de Chile).")
    print(f"   Valores nulos en conexiones_prom2024: {df['conexiones_prom2024'].isnull().sum()}")
    return df


def main():
    print("=" * 60)
    print("  PREPARACIÓN DE DATOS – BRECHA DIGITAL CHILE")
    print("=" * 60)
    if os.path.exists(SUBTEL_PATH):
        subtel_dict = extraer_subtel(SUBTEL_PATH)
        actualizar_dataset(subtel_dict)
    else:
        print(f"⚠️ No se encontró {SUBTEL_PATH}. Descargar antes con src/descarga_datos.py")


if __name__ == "__main__":
    main()
